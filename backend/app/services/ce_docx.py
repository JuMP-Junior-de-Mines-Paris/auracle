# app/services/ce_docx.py
"""Remplit le modèle Word de la convention d'étude.

On part du vrai document JuMP (app/templates/convention_etude.docx) plutôt
que d'en regénérer un : mise en page, en-têtes, logos et conditions
générales restent ceux de la Junior-Entreprise, et une retouche du modèle
dans Word se retrouve dans toutes les conventions suivantes.

Trois sortes d'opérations, dans cet ordre :
  1. structurelles — dupliquer les phases, les points de contrôle, les
     lignes des tableaux (budget, planning) selon les données ;
  2. ciblées — les phrases du modèle qui ne sont pas des {placeholders}
     (« [X mois] », « janvier 2026 », « %Numéro% année »…) ;
  3. génériques — les {placeholders}, y compris ceux que Word a éclatés
     sur plusieurs « runs » (« {», « ref_CE », « } »).
"""
from __future__ import annotations

import copy
import io
import re
import unicodedata
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from app.services.ce_budget import (
    Computed, compute, date_fr, eur_en_lettres, fmt_eur, fmt_pct,
)
from app.services.ce_schema import CEData

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "convention_etude.docx"

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


# ─── Primitives XML ────────────────────────────────────────────────────────

def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _norm(key: str) -> str:
    key = unicodedata.normalize("NFD", key.strip())
    return key.encode("ascii", "ignore").decode().lower()


def _set_t(t, value: str) -> None:
    t.text = value
    t.set(XML_SPACE, "preserve")


# Les accolades saisies par l'utilisateur ne doivent jamais être prises pour
# des {placeholders} du modèle : « {total_ht} » écrit dans un cahier des
# charges serait sinon remplacé par un montant. On les met à l'abri pendant
# le remplissage et on les rétablit à la fin (_unshield).
_SHIELD = str.maketrans({"{": "\ue001", "}": "\ue002"})
_UNSHIELD = str.maketrans({"\ue001": "{", "\ue002": "}"})


def _shield(value: str) -> str:
    return value.translate(_SHIELD)


def _unshield(doc) -> None:
    for p in _all_paragraphs(doc):          # corps, en-têtes et pieds de page
        for t in p.iter(qn("w:t")):
            if t.text and ("\ue001" in t.text or "\ue002" in t.text):
                t.text = t.text.translate(_UNSHIELD)


def sub_par(p, pattern: str, repl) -> int:
    """Remplace `pattern` dans un paragraphe, même éclaté sur plusieurs runs.

    La mise en forme du premier run touché est conservée ; les runs
    suivants perdent la portion remplacée (et deviennent vides au besoin).
    `repl` est une chaîne ou une fonction de l'objet match.
    """
    ts = list(p.iter(qn("w:t")))
    originals = [t.text or "" for t in ts]
    texts = list(originals)
    full = "".join(originals)
    matches = list(re.finditer(pattern, full))
    if not matches:
        return 0

    starts, pos = [], 0
    for o in originals:
        starts.append(pos)
        pos += len(o)

    # En partant de la fin : les décalages des correspondances précédentes
    # restent valables.
    for m in reversed(matches):
        s, e = m.span()
        new = _shield(repl(m) if callable(repl) else repl)
        first = True
        for i, o in enumerate(originals):
            seg_s, seg_e = starts[i], starts[i] + len(o)
            if seg_e <= s or seg_s >= e:
                continue
            lo, hi = max(s, seg_s) - seg_s, min(e, seg_e) - seg_s
            texts[i] = texts[i][:lo] + (new if first else "") + texts[i][hi:]
            first = False
    for t, new_text in zip(ts, texts):
        if t.text != new_text:
            _set_t(t, new_text)
    return len(matches)


def set_par_text(p, value: str) -> None:
    """Remplace tout le texte d'un paragraphe en gardant la mise en forme du premier run."""
    ts = list(p.iter(qn("w:t")))
    if not ts:
        return
    _set_t(ts[0], _shield(value))
    for t in ts[1:]:
        _set_t(t, "")


def _clone(el):
    """Copie profonde, sans les identifiants de paragraphe que Word veut uniques."""
    c = copy.deepcopy(el)
    for node in c.iter():
        for attr in (f"{{{W14}}}paraId", f"{{{W14}}}textId"):
            node.attrib.pop(attr, None)
    return c


def _remove(el) -> None:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _body_paragraphs(doc):
    return [el for el in doc.element.body.iter(qn("w:p"))]


def find_par(doc, pattern: str):
    rx = re.compile(pattern)
    for p in _body_paragraphs(doc):
        if rx.search(_text(p)):
            return p
    return None


def find_table(doc, needle: str):
    for tbl in doc.element.body.iter(qn("w:tbl")):
        if needle in _text(tbl):
            return tbl
    raise LookupError(f"Tableau introuvable dans le modèle : {needle!r}")


def _cell_set(tc, value: str) -> None:
    paras = tc.findall(qn("w:p"))
    set_par_text(paras[0], value)
    for extra in paras[1:]:
        _remove(extra)


def _next_siblings(el, n: int) -> list:
    out, cur = [el], el
    for _ in range(n - 1):
        cur = cur.getnext()
        out.append(cur)
    return out


def fill_prose(p, text: str) -> None:
    """Un paragraphe-marqueur devient autant de paragraphes que de lignes."""
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    lines = [("• " + ln[2:]) if ln.startswith(("- ", "* ")) else ln for ln in lines]
    if not lines:
        set_par_text(p, "")
        return
    set_par_text(p, lines[0])
    cursor = p
    for ln in lines[1:]:
        clone = _clone(p)
        set_par_text(clone, ln)
        cursor.addnext(clone)
        cursor = clone


# ─── 0. Notes de rédaction du modèle ───────────────────────────────────────

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _strip_editorial_notes(doc) -> int:
    """Retire les encadrés « Comment faire – À RETIRER À LA FIN ».

    Le modèle en contient une trentaine : ce sont des consignes à l'adresse
    du rédacteur, pas du client. Chacun est une zone de texte dans un run.
    """
    n = 0
    for run in list(doc.element.body.iter(qn("w:r"))):
        holds_box = (
            run.find(f"{{{MC}}}AlternateContent") is not None
            or run.find(qn("w:drawing")) is not None
            or run.find(qn("w:pict")) is not None
        )
        if holds_box and "Comment faire" in _text(run):
            _remove(run)
            n += 1
        elif _is_note_screenshot(run):
            _remove(run)
            n += 1
    return n


def _is_note_screenshot(run) -> bool:
    """Une des consignes est collée dans le modèle sous forme d'image (« Zone de texte »)."""
    for doc_pr in run.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
        if (doc_pr.get("descr") or "").strip() == "Zone de texte":
            return True
    return False


LITIGES = {
    "administration_publique": (
        "En cas de désaccord persistant, le litige sera porté devant le tribunal "
        "administratif dont dépend le siège de JuMP."
    ),
    "junior_sous_traitance": (
        "En cas de désaccord persistant, le litige sera porté en premier lieu devant le "
        "Conseil d’Administration de la Confédération Nationale des Junior-Entreprises. "
        "Dans le cas où le litige ne serait pas résolu, celui-ci sera porté devant le "
        "tribunal judiciaire dont dépend le siège de JuMP."
    ),
}


def _litiges_clause(doc, data: CEData) -> None:
    """Article 13 : le tribunal compétent dépend du type de client."""
    variant = LITIGES.get(data.client.type_client)
    if variant is None:
        return
    p = find_par(doc, r"^La présente Convention est soumise au droit français")
    if p is not None:
        sub_par(p, r"En cas de désaccord persistant.*$", variant)


# ─── 1. Opérations structurelles ──────────────────────────────────────────

def _fill_prose_blocks(doc, data: CEData) -> None:
    for key, value in {
        "etude_contexte_et_problématique": data.contexte,
        "etude_existant": data.existant,
        "etude_resumecahierdescharges": data.resume_cdc,
        "etude_cahierdescharges": data.cdc,
    }.items():
        p = find_par(doc, r"^\s*\{\s*" + re.escape(key) + r"\s*\}\s*$")
        if p is not None:
            fill_prose(p, value)

    exemples = find_par(doc, r"^\s*\{exemples\}\s*$")
    if exemples is not None:
        if data.exemples.strip():
            fill_prose(exemples, data.exemples)
        else:
            # Sans exemple, la phrase d'introduction n'a plus d'objet.
            intro = find_par(doc, r"nous avons déjà effectué des études")
            _remove(exemples)
            if intro is not None:
                _remove(intro)


def _fill_phases(doc, data: CEData, c: Computed) -> None:
    proto_head = find_par(doc, r"^Phase 1 :")
    group = _next_siblings(proto_head, 4)            # titre, blanc, description, blanc
    x_head = find_par(doc, r"^Phase\s*X\s*:")
    for el in _next_siblings(x_head, 4):
        _remove(el)

    anchor = group[-1]
    for ph, line in zip(data.phases, c.lines):
        block = [_clone(el) for el in group]
        set_par_text(block[0], f"Phase {line.index} : {ph.nom} - {ph.jeh} JEH")
        for el in block:
            anchor.addnext(el)
            anchor = el
        # Une fois dans le document : une description multi-lignes y ajoute des paragraphes.
        fill_prose(block[2], ph.description)
        anchor = block[3]
    for el in group:
        _remove(el)


def _fill_checkpoints(doc, data: CEData) -> None:
    first = find_par(doc, r"^Point de contrôle 1")
    group = _next_siblings(first, 4)                 # titre, format, requis, blanc
    x_first = find_par(doc, r"^Point de contrôle\s*X")
    for el in _next_siblings(x_first, 3):
        _remove(el)

    anchor = group[-1]
    for i, pc in enumerate(data.points_controle, start=1):
        titre, fmt, requis, blanc = (_clone(el) for el in group)
        set_par_text(titre, f"Point de contrôle {i} :" + (f" {pc.nom}" if pc.nom else ""))
        set_par_text(fmt, f"Format : {pc.format}")
        if pc.requis_avant_phase is None:
            requis = None
        else:
            set_par_text(requis, f"Nécessaire au début de la phase : {pc.requis_avant_phase}")
        for el in (titre, fmt, requis, blanc):
            if el is not None:
                anchor.addnext(el)
                anchor = el
    for el in group:
        _remove(el)


def _fill_planning(doc, data: CEData, c: Computed) -> None:
    tbl = find_table(doc, "Durée (semaines)")
    rows = tbl.findall(qn("w:tr"))
    head0, head1, proto = rows[0], rows[1], rows[2]
    for r in rows[3:]:
        _remove(r)

    n = c.semaines
    week_area = 7 * 881                              # largeur occupée par les semaines dans le modèle
    width = max(week_area // n, 300)
    compact = n > 7                                  # « Semaine » ne tient plus : « S3 »

    filled = _clone(proto.findall(qn("w:tc"))[3])
    empty = _clone(proto.findall(qn("w:tc"))[5])

    def _set_width(tc):
        w = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
        w.set(qn("w:w"), str(width))
        w.set(qn("w:type"), "dxa")

    def _rebuild(row, week_proto, text_fn=None):
        cells = row.findall(qn("w:tc"))
        for tc in cells[3:]:
            _remove(tc)
        for w in range(1, n + 1):
            tc = _clone(week_proto(w))
            _set_width(tc)
            if text_fn:
                text_fn(tc, w)
            row.append(tc)

    h0 = _clone(head0.findall(qn("w:tc"))[3])
    h1 = _clone(head1.findall(qn("w:tc"))[3])

    def _head_text(tc, w):
        # « Semaine » et le numéro sont dans deux w:t (parfois deux paragraphes).
        ts = list(tc.iter(qn("w:t")))
        labels = [f"S{w}"] if compact else ["Semaine", str(w)]
        for i, t in enumerate(ts):
            _set_t(t, labels[i] if i < len(labels) else "")

    _rebuild(head0, lambda w: h0, _head_text)
    _rebuild(head1, lambda w: h1)

    for line in c.lines:
        row = _clone(proto)
        cells = row.findall(qn("w:tc"))
        _cell_set(cells[0], f"{line.index} : {line.nom}")
        _cell_set(cells[1], str(line.duree))
        _cell_set(cells[2], str(line.jeh))
        _rebuild(row, lambda w, l=line: filled if l.debut <= w <= l.fin else empty)
        tbl.append(row)
    _remove(proto)

    grid = tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    for g in cols[3:]:
        grid.remove(g)
    for _ in range(n):
        g = copy.deepcopy(cols[3])
        g.set(qn("w:w"), str(width))
        grid.append(g)


def _fill_budget_table(doc, data: CEData, c: Computed) -> None:
    tbl = find_table(doc, "Prix H.T. / JEH")
    rows = tbl.findall(qn("w:tr"))
    proto_title, proto_total = rows[1], rows[2]
    totals_start = rows[7]                           # « Total H.T. hors frais »
    for r in rows[3:7]:
        _remove(r)
    # Pas de frais de structure : le total est la somme des phases. On retire
    # la ligne des frais et son montant en lettres, et la ligne de somme
    # (nombre de JEH + montant) devient « Total des phases ».
    for r in rows[8:10]:
        _remove(r)
    _cell_set(totals_start.findall(qn("w:tc"))[0], "Total des phases")

    for line in c.lines:
        title, total = _clone(proto_title), _clone(proto_total)
        _cell_set(title.findall(qn("w:tc"))[0], f"Phase {line.index} : {line.nom}")
        cells = total.findall(qn("w:tc"))
        _cell_set(cells[0], f"Total phase {line.index}")
        _cell_set(cells[1], str(line.jeh))
        _cell_set(cells[2], f"{fmt_eur(line.prix_jeh)} € HT")
        _cell_set(cells[3], f"{fmt_eur(line.total_ht)} € HT")
        totals_start.addprevious(title)
        totals_start.addprevious(total)
    _remove(proto_title)
    _remove(proto_total)


def _fill_payment_table(doc, data: CEData, c: Computed) -> None:
    tbl = find_table(doc, "Modalités de paiement")
    rows = tbl.findall(qn("w:tr"))
    pay = data.paiement
    has_inter = pay.intermediaire_pct > 0

    if has_inter:
        sub_par(rows[0].findall(qn("w:tc"))[2], r"XXX",
                pay.livrable_intermediaire.strip() or "intermédiaire")
        _cell_set(rows[2].findall(qn("w:tc"))[2], f"{fmt_eur(c.intermediaire_ht)} €")
        _cell_set(rows[3].findall(qn("w:tc"))[2], f"{fmt_eur(c.intermediaire_ttc)} €")
    else:
        # Pas de versement intermédiaire : la colonne entière disparaît et
        # les deux autres se répartissent sa largeur.
        for r in rows:
            _remove(r.findall(qn("w:tc"))[2])
        grid = tbl.find(qn("w:tblGrid"))
        cols = grid.findall(qn("w:gridCol"))
        total = sum(int(g.get(qn("w:w"))) for g in cols)
        grid.remove(cols[2])
        left = grid.findall(qn("w:gridCol"))
        share = total // len(left)
        for g in left:
            g.set(qn("w:w"), str(share))
        for r in rows:
            for tc in r.findall(qn("w:tc")):
                w = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
                if w is not None:
                    w.set(qn("w:w"), str(share))
                    w.set(qn("w:type"), "dxa")


def _fill_team(doc, data: CEData) -> None:
    tbl = find_table(doc, "Nom Prénom")
    cdp = data.chef_projet
    rows = tbl.findall(qn("w:tr"))
    _cell_set(rows[0].findall(qn("w:tc"))[0], cdp.nom_complet)
    _cell_set(rows[1].findall(qn("w:tc"))[0], cdp.formation)


def _fee_clause(doc, data: CEData) -> None:
    """Article 5 : on garde la clause choisie, on retire l'autre et les titres entre crochets."""
    keep, drop = (("réel", "forfait") if data.budget.frais_mode == "reel" else ("forfait", "réel"))
    head = find_par(doc, rf"^\[Facturation des frais au {drop}\]")
    if head is not None:
        _remove(head.getnext())
        _remove(head)
    head = find_par(doc, rf"^\[Facturation des frais au {keep}\]")
    if head is not None:
        _remove(head)


# ─── 2. Phrases du modèle qui ne sont pas des {placeholders} ───────────────

def _targeted(doc, data: CEData, c: Computed) -> None:
    p = find_par(doc, r"^L’étude prendra fin")
    if p is not None:
        sub_par(p, r"\[X\]", str(c.semaines))

    p = find_par(doc, r"Cette convention d’étude est valable")
    if p is not None:
        sub_par(p, r"\[X mois\]", f"{data.validite_mois} mois")

    p = find_par(doc, r"Si l’étude se poursuit au-delà de")
    if p is not None:
        fin = data.chef_projet.fin_mandat.strip()
        if fin:
            sub_par(p, r"janvier 2026", fin)
        else:
            _remove(p)

    n = len(data.intervenants)
    p = find_par(doc, r"^L’intervenant")
    if p is not None:
        sub_par(p, r"\s*Mettre au pluriel si plusieurs intervenants", "")
        if n > 1:
            sub_par(p, r"^L’intervenant", "Les intervenants")

    p = find_par(doc, r"propose de confier l'étude à un étudiant")
    if p is not None:
        def _un(iv):
            qual = f" et {iv.qualification.strip()}" if iv.qualification.strip() else ""
            etudiant = "étudiante" if iv.titre == "Madame" else "étudiant"
            nom = f"{iv.nom_complet}, " if iv.nom_complet else "un "
            return f"{nom}{etudiant} de {iv.annee.strip()} année de l’École{qual}"
        if n == 0:
            texte = "un étudiant de l’École"
        elif n == 1:
            texte = _un(data.intervenants[0])
        else:
            texte = f"{n} étudiants : " + " ; ".join(_un(iv) for iv in data.intervenants)
        sub_par(p, r"un étudiant de %Numéro%.*?Option_double_diplôme%", texte)

    p = find_par(doc, r"Pour \[Nom Prénom et adresse du Chargé d’étude\]")
    if p is not None:
        cdp = data.chef_projet
        sub_par(p, r"\[Nom Prénom et adresse du Chargé d’étude\]\s*",
                ", ".join(x for x in (cdp.nom_complet, cdp.email) if x))
    p = find_par(doc, r"Pour \[Nom Prénom et adresse du Président\]")
    if p is not None:
        pres = data.president
        sub_par(p, r"\[Nom Prénom et adresse du Président\]\s*",
                ", ".join(x for x in (pres.nom_complet, pres.email) if x))

    p = find_par(doc, r"Document comportant X pages")
    if p is not None:
        _page_count_field(p)


def _page_count_field(p) -> None:
    """« X pages » devient un champ NUMPAGES : Word et LibreOffice le calculent."""
    marker = "\ue000"
    sub_par(p, r"comportant X pages", f"comportant {marker} pages")
    for r in p.iter(qn("w:r")):
        t = r.find(qn("w:t"))
        if t is None or marker not in (t.text or ""):
            continue
        before, after = t.text.split(marker, 1)
        _set_t(t, before)

        def _run_with(child):
            run = etree.Element(qn("w:r"))
            props = r.find(qn("w:rPr"))
            if props is not None:
                run.append(copy.deepcopy(props))
            run.append(child)
            return run

        def _fld(kind):
            e = etree.Element(qn("w:fldChar"))
            e.set(qn("w:fldCharType"), kind)
            return e

        instr = etree.Element(qn("w:instrText"))
        instr.set(XML_SPACE, "preserve")
        instr.text = " NUMPAGES "
        cached = etree.Element(qn("w:t"))
        cached.text = "30"
        tail = etree.Element(qn("w:t"))
        _set_t(tail, after)

        cursor = r
        for child in (_fld("begin"), instr, _fld("separate"), cached, _fld("end"), tail):
            run = _run_with(child)
            cursor.addnext(run)
            cursor = run
        return


# ─── 3. Placeholders génériques ────────────────────────────────────────────

def _mapping(data: CEData, c: Computed) -> tuple[list[tuple[str, str]], dict[str, str]]:
    cl = data.client
    rep = cl.representant
    contact = cl.interlocuteur or rep
    pres = data.president

    # Montants en lettres : le modèle écrit « {x_lettres} euros HT » ; le mot
    # « euros » est dans notre phrase (« un euro », « et cinquante centimes »).
    phrases = [
        (r"(?i)\{total_ht_lettres\}\s*euros", eur_en_lettres(c.total_ht)),
        (r"(?i)\{total_ttc_lettres\}\s*euros", eur_en_lettres(c.total_ttc)),
        # SIRET : le modèle ne porte pas le mot, et certains clients n'en ont pas.
        (r",\s*\{SIRET\}\s*", f", SIRET {cl.siret.strip()} " if cl.siret.strip() else " "),
        (r"\{président_titre[^}]*\}", pres.nom_complet),
    ]

    simple = {
        "client_titre": contact.titre,
        "client_prenom": contact.prenom,
        "client_nom": contact.nom,
        "ref_ce": data.ref_ce,
        "num_ap": data.ref_ce,
        "nom_societe": cl.societe,
        "reference_etude": data.reference_etude,
        "etude_client_societe": cl.societe,
        "etude_client_societe_nom": cl.societe,
        "etude_client_societe_adresse": cl.adresse,
        "etude_client_societe_code_postal": cl.code_postal,
        "etude_client_societe_ville": cl.ville,
        "etude_client_societe_pays": cl.pays,
        "etude_client_titre": rep.titre,
        "etude_client_prenom": rep.prenom,
        "etude_client_nom": rep.nom,
        "etude_client_nom_upper": rep.nom.upper(),
        "etude_client_fonction": rep.fonction,
        "client_interlocuteur_nom": contact.nom_complet,
        "client_interlocuteur_mail": contact.email,
        "je_president_titre": pres.titre,
        "je_president_prenom": pres.prenom,
        "je_president_nom": pres.nom,
        "je_president_fonction": pres.fonction or "Président",
        "date_redaction": date_fr(data.date_redaction),
        "nb_jehs": str(c.nb_jeh),
        "prix_ht": fmt_eur(c.prix_ht),
        "total_ht": fmt_eur(c.total_ht),
        "taux_tva": fmt_pct(c.taux_tva),
        "tva_sur_total": fmt_eur(c.tva),
        "total_ttc": fmt_eur(c.total_ttc),
        "etude_acompte_ht": fmt_eur(c.acompte_ht),
        "etude_acompte_ttc": fmt_eur(c.acompte_ttc),
        "etude_reste_ht": fmt_eur(c.reste_ht),
        "etude_reste_ttc": fmt_eur(c.reste_ttc),
    }
    return phrases, simple


def _all_paragraphs(doc):
    """Corps, tableaux, puis en-têtes et pieds de page de chaque section."""
    yield from doc.element.body.iter(qn("w:p"))
    seen = set()
    for section in doc.sections:
        for part in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            el = part._element
            if id(el) in seen:
                continue
            seen.add(id(el))
            yield from el.iter(qn("w:p"))


def _generic(doc, data: CEData, c: Computed) -> None:
    phrases, simple = _mapping(data, c)
    placeholder = re.compile(r"\{([^{}]+)\}")

    def lookup(m):
        value = simple.get(_norm(m.group(1)))
        return m.group(0) if value is None else value

    for p in _all_paragraphs(doc):
        txt = _text(p)
        if "{" not in txt:
            continue
        for pattern, value in phrases:
            sub_par(p, pattern, value)
        sub_par(p, placeholder.pattern, lookup)


# ─── Mise en page ──────────────────────────────────────────────────────────

SECTION_TITLES = (
    "SOMMAIRE", "PRÉSENTATION DE JUMP", "REFERENCES", "CONTEXTE, PROBLEMATIQUE",
    "DÉFINITION DE L’ÉTUDE", "CAHIER DES CHARGES", "MÉTHODOLOGIE DE L’ETUDE", "PLANNING",
    "BUDGET ET CONDITIONS", "PRÉSENTATION DE L’ÉQUIPE", "CONDITIONS GENERALES", "CONCLUSION",
)
TITLE_STYLES = {"Titre-Middle", "Heading 1"}


def _is_filler(el) -> bool:
    """Paragraphe vide qui ne sert qu'à pousser le contenu (ni image, ni saut, ni champ)."""
    if el.tag != qn("w:p") or _text(el).strip():
        return False
    return not any(
        el.find(f".//{tag}") is not None
        for tag in (qn("w:drawing"), qn("w:pict"), qn("w:br"), qn("w:sectPr"),
                    qn("w:fldChar"), f"{{{MC}}}AlternateContent")
    )


def _breaks_page(el) -> bool:
    if el is None:
        return False
    if el.find(f".//{qn('w:sectPr')}") is not None:
        return True
    return any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br")))


def _paginate_sections(doc) -> None:
    """Chaque grande section démarre en haut d'une page.

    Le modèle obtenait cet effet avec des paragraphes vides dimensionnés
    pour la longueur de ses encadrés de consigne ; une fois ceux-ci retirés,
    les titres atterrissent en bas de page. On remplace le remplissage par
    un vrai saut de page.
    """
    from docx.text.paragraph import Paragraph

    for p in [el for el in doc.element.body if el.tag == qn("w:p")]:
        para = Paragraph(p, doc._body)
        if para.style.name not in TITLE_STYLES:
            continue
        label = _text(p).strip().upper()
        if not any(label.startswith(t) for t in SECTION_TITLES):
            continue
        prev = p.getprevious()
        while prev is not None and _is_filler(prev):
            before = prev.getprevious()
            _remove(prev)
            prev = before
        if not _breaks_page(prev):
            para.paragraph_format.page_break_before = True


# Le modèle écrit en rouge tout ce qui reste à remplir. Une fois rempli, le
# texte reprend la couleur du document.
RED = {"FF0000", "EE0000"}


def _recolor_filled_text(doc) -> None:
    roots = [doc.element.body]
    for section in doc.sections:
        roots += [section.header._element, section.footer._element,
                  section.first_page_header._element, section.first_page_footer._element,
                  section.even_page_header._element, section.even_page_footer._element]
    for root in roots:
        for color in root.iter(qn("w:color")):
            if (color.get(qn("w:val")) or "").upper() in RED:
                color.set(qn("w:val"), "000000")


def _refresh_fields_on_open(doc) -> None:
    """Sommaire et numéros de page : Word les recalcule à l'ouverture.

    L'ordre des éléments de settings.xml est imposé par le schéma, et Word
    est plus strict que LibreOffice : updateFields doit précéder
    hdrShapeDefaults (ou, à défaut, les notes de bas de page et la compat).
    """
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    el = etree.Element(qn("w:updateFields"))
    el.set(qn("w:val"), "true")
    for successor in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids"):
        anchor = settings.find(qn(successor))
        if anchor is not None:
            anchor.addprevious(el)
            return
    settings.append(el)


def _set_metadata(doc, data: CEData) -> None:
    """Le modèle porte les noms de ses anciens rédacteurs : le document est celui de JuMP."""
    props = doc.core_properties
    props.title = f"Convention d'Étude {data.ref_ce}"
    props.subject = data.reference_etude
    props.author = "JuMP"
    props.last_modified_by = "JuMP Oracle"
    props.comments = ""
    props.keywords = ""


# ─── Point d'entrée ────────────────────────────────────────────────────────

def render_ce(data: CEData, template: Path = TEMPLATE) -> bytes:
    doc = Document(str(template))
    c = compute(data)

    _strip_editorial_notes(doc)
    _fill_prose_blocks(doc, data)
    _fill_phases(doc, data, c)
    _fill_checkpoints(doc, data)
    _fill_planning(doc, data, c)
    _fill_budget_table(doc, data, c)
    _fill_payment_table(doc, data, c)
    _fill_team(doc, data)
    _fee_clause(doc, data)
    _litiges_clause(doc, data)
    _targeted(doc, data, c)
    _generic(doc, data, c)
    _unshield(doc)
    _paginate_sections(doc)
    _recolor_filled_text(doc)
    _refresh_fields_on_open(doc)
    _set_metadata(doc, data)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
