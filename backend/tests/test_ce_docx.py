import io
import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.services.ce_docx import render_ce
from app.services.ce_schema import CEData
from tests.ce_sample import sample


def _render(**overrides):
    d = sample()
    for path, value in overrides.items():
        node = d
        *parents, last = path.split("__")
        for key in parents:
            node = node[int(key)] if key.isdigit() else node[key]
        node[int(last) if last.isdigit() else last] = value
    return Document(io.BytesIO(render_ce(CEData.model_validate(d))))


def _paragraphs(doc):
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def _text(doc):
    return "\n".join(_paragraphs(doc))


def _table(doc, needle):
    for tbl in doc.element.body.iter(qn("w:tbl")):
        if needle in "".join(t.text or "" for t in tbl.iter(qn("w:t"))):
            return tbl
    raise AssertionError(needle)


def _rows(tbl):
    return [
        ["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.findall(qn("w:tc"))]
        for tr in tbl.findall(qn("w:tr"))
    ]


@pytest.fixture(scope="module")
def doc():
    return _render()


def test_no_placeholder_or_template_instruction_survives(doc):
    text = _text(doc)
    assert not re.search(r"\{[^}]*\}", text)
    assert not re.search(r"\[[^\]]*\]", text)
    assert "XXX" not in text and "%Numéro%" not in text
    assert "Mettre au pluriel" not in text


def test_editorial_notes_are_gone(doc):
    text = _text(doc)
    assert "Comment faire" not in text and "RETIRER" not in text
    xml = doc.element.body.xml
    assert 'descr="Zone de texte"' not in xml
    assert "AlternateContent" not in xml


def test_footer_placeholders_are_filled(doc):
    footers = " ".join(
        "".join(t.text or "" for t in s.footer._element.iter(qn("w:t"))) for s in doc.sections
    )
    assert "Acme Recyclage SAS" in footers and "{" not in footers


def test_cover_and_parties(doc):
    text = _text(doc)
    assert "À l’attention de Monsieur Paul Martin" in text
    assert "CE-2026-014" in text
    assert "12 octobre 2026" in text
    assert "SIRET 123 456 789 00012" in text
    assert "Directrice générale" in text


def test_phases_are_cloned_in_order_with_their_descriptions(doc):
    paras = _paragraphs(doc)
    heads = [p for p in paras if re.match(r"Phase \d : ", p) and " - " in p]
    assert heads == ["Phase 1 : Cadrage - 3 JEH", "Phase 2 : Modélisation - 8 JEH", "Phase 3 : Restitution - 2 JEH"]
    i = paras.index("Phase 2 : Modélisation - 8 JEH")
    # Une description sur deux lignes devient deux paragraphes.
    assert "Construction du modèle technico-économique." in paras[i:i + 5]
    assert "Validation des hypothèses." in paras[i:i + 5]


def test_prose_blocks_are_split_into_paragraphs(doc):
    paras = _paragraphs(doc)
    assert "Acme souhaite valoriser ses déchets de batteries." in paras
    assert "Elle cherche un procédé rentable." in paras
    assert "• Revue de la littérature" in paras


def test_checkpoints(doc):
    text = _text(doc)
    assert "Point de contrôle 1 : Rapport intermédiaire" in text
    assert "Nécessaire au début de la phase : 3" in text
    assert "Point de contrôle 2 : Rapport final" in text
    # Le second n'a pas de phase à débloquer : la ligne n'est pas écrite.
    assert text.count("Nécessaire au début de la phase") == 1


def test_budget_table(doc):
    rows = _rows(_table(doc, "Prix H.T. / JEH"))
    flat = [" | ".join(r) for r in rows]
    assert "Phase 2 : Modélisation" in flat[3]
    assert flat[2].startswith("Total phase 1")
    joined = "\n".join(flat)
    assert "13 JEHs" in joined
    assert "1 566,50 € HT" in joined
    assert "Total des phases" in joined
    assert "Mille cinq cent soixante-six euros et cinquante centimes HT" in joined
    assert "313,30 €" in joined
    assert "1 879,80 € TTC" in joined
    assert "Mille huit cent soixante-dix-neuf euros et quatre-vingts centimes TTC" in joined



def test_budget_has_one_pair_of_rows_per_phase():
    doc = _render(points_controle=[], phases=[
        {"nom": f"P{i}", "jeh": 1, "duree_semaines": 1, "description": "d"} for i in range(1, 6)
    ])
    rows = _rows(_table(doc, "Prix H.T. / JEH"))
    assert sum(1 for r in rows if r[0].startswith("Total phase")) == 5


def test_planning_has_one_column_per_week_and_shades_the_right_ones(doc):
    tbl = _table(doc, "Durée (semaines)")
    rows = tbl.findall(qn("w:tr"))
    assert len(rows) == 5                                   # 2 en-têtes + 3 phases
    assert len(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))) == 3 + 7

    def shaded(tr):
        cells = tr.findall(qn("w:tc"))[3:]
        return [
            c.find(qn("w:tcPr")).find(qn("w:shd")) is not None
            and c.find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill")) == "564156"
            for c in cells
        ]
    assert shaded(rows[2]) == [True, True, False, False, False, False, False]
    assert shaded(rows[3]) == [False, False, True, True, True, True, False]
    assert shaded(rows[4]) == [False] * 6 + [True]
    assert "L’étude prendra fin 7 semaines" in _text(doc)


def test_long_studies_get_compact_week_headers():
    doc = _render(points_controle=[], phases=[{"nom": "Long", "jeh": 10, "duree_semaines": 12, "description": "d"}])
    tbl = _table(doc, "Durée (semaines)")
    assert len(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))) == 3 + 12
    header = _rows(tbl)[0]
    assert header[3:] == [f"S{i}" for i in range(1, 13)]


def test_payment_table_without_intermediate_drops_the_column(doc):
    rows = _rows(_table(doc, "Modalités de paiement"))
    assert all(len(r) == 3 for r in rows)
    assert rows[2][1:] == ["469,95 €", "1 096,55 €"]
    assert rows[3][1:] == ["563,94 €", "1 315,86 €"]


def test_payment_table_with_intermediate_keeps_three_installments():
    doc = _render(paiement={"acompte_pct": "30", "intermediaire_pct": "30", "livrable_intermediaire": "Rapport intermédiaire"})
    rows = _rows(_table(doc, "Modalités de paiement"))
    assert all(len(r) == 4 for r in rows)
    assert "Rapport intermédiaire" in rows[0][2]
    assert rows[2][2] == "469,95 €"
    assert "XXX" not in _text(doc)


def test_fee_clause_follows_the_chosen_mode(doc):
    text = _text(doc)
    assert "refacturés au réel" in text and "facturation au forfait" not in text
    forfait = _text(_render(budget__frais_mode="forfait"))
    assert "facturation au forfait" in forfait and "refacturés au réel" not in forfait.split("Article 5")[1]
    assert "[Facturation" not in forfait


def test_client_type_changes_the_jurisdiction_clause():
    default = _text(_render())
    assert "tribunal judiciaire dont dépend le siège de JuMP" in default
    admin = _text(_render(client__type_client="administration_publique"))
    assert "tribunal administratif" in admin
    sous = _text(_render(client__type_client="junior_sous_traitance"))
    assert "Conseil d’Administration de la Confédération Nationale" in sous


def test_team_and_intervenants(doc):
    text = _text(doc)
    assert "Léa Toledano" in text
    assert "Étudiante en deuxième année" in text
    assert "au-delà de janvier 2027" in text
    assert "Camille Durand, étudiante de deuxième année de l’École et en double diplôme avec l'ENSAE" in text
    assert "L’intervenant" in text and "Les intervenants" not in text


def test_several_intervenants_pluralise():
    doc = _render(intervenants=[
        {"prenom": "Camille", "nom": "Durand", "titre": "Madame", "annee": "deuxième", "qualification": ""},
        {"prenom": "Paul", "nom": "Roux", "annee": "troisième", "qualification": ""},
    ])
    text = _text(doc)
    assert "Les intervenants" in text
    assert "2 étudiants : Camille Durand, étudiante de deuxième année de l’École ; Paul Roux, étudiant de troisième année de l’École" in text


def test_optional_sections_disappear_cleanly():
    no_examples = _text(_render(exemples=""))
    assert "nous avons déjà effectué des études" not in no_examples
    no_end = _text(_render(chef_projet={**sample()["chef_projet"], "fin_mandat": ""}))
    assert "nouveau chef de projet" not in no_end


def test_missing_siret_leaves_no_dangling_punctuation():
    d = sample()
    d["client"]["siret"] = ""
    text = _text(Document(io.BytesIO(render_ce(CEData.model_validate(d)))))
    assert "SIRET" not in text.split("ET\n")[-1].split("représentée par")[0]
    assert ", ," not in text and ",  " not in text


def test_validity_and_page_count_field(doc):
    assert "valable 3 mois" in _text(doc)
    assert "NUMPAGES" in doc.element.body.xml


def test_to_fill_red_is_gone(doc):
    colors = {c.get(qn("w:val")).upper() for c in doc.element.body.iter(qn("w:color"))}
    assert not colors & {"FF0000", "EE0000"}


def test_sections_start_on_a_new_page(doc):
    # Les consignes retirées servaient de remplissage : sans saut explicite,
    # ces titres tomberaient en bas de page.
    def breaks_before(title):
        for p in doc.element.body.iter(qn("w:p")):
            if "".join(t.text or "" for t in p.iter(qn("w:t"))).strip() == title:
                return p.find(qn("w:pPr")).find(qn("w:pageBreakBefore")) is not None
        raise AssertionError(title)
    for title in ("PLANNING", "DÉFINITION DE L’ÉTUDE", "MÉTHODOLOGIE DE L’ETUDE"):
        assert breaks_before(title), title


def test_toc_is_refreshed_when_word_opens_the_file(doc):
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_text_with_xml_special_characters_is_escaped():
    doc = _render(contexte="Chiffre d'affaires < 5 M€ & marge > 10 % — « test »")
    assert "Chiffre d'affaires < 5 M€ & marge > 10 % — « test »" in _text(doc)


def test_update_fields_sits_where_the_schema_puts_it(doc):
    # Word refuse (ou « répare ») un settings.xml dont l'ordre viole le schéma.
    children = [c.tag.split("}")[1] for c in doc.settings.element]
    assert children.index("updateFields") < children.index("hdrShapeDefaults")
    assert children.count("updateFields") == 1


def test_metadata_is_jumps_not_the_template_authors(doc):
    props = doc.core_properties
    assert props.title == "Convention d'Étude CE-2026-014"
    assert props.author == "JuMP"
    assert "Rinchet" not in (props.last_modified_by or "") and "Beaurain" not in (props.author or "")


def test_a_minimal_draft_renders_without_leftovers():
    minimal = {
        "ref_ce": "CE-1", "reference_etude": "Étude", "date_redaction": "2026-01-01",
        "client": {"societe": "Zeta", "representant": {"prenom": "A", "nom": "B"}},
        "president": {"prenom": "C", "nom": "D"}, "chef_projet": {"prenom": "E", "nom": "F"},
        "phases": [{"nom": "Unique", "jeh": 1, "duree_semaines": 1}],
        "budget": {"prix_jeh": "80"},
        "paiement": {"acompte_pct": "50", "intermediaire_pct": "25", "livrable_intermediaire": "V1"},
    }
    text = _text(Document(io.BytesIO(render_ce(CEData.model_validate(minimal)))))
    assert not re.search(r"\{[^}]*\}|\[[^\]]*\]|XXX", text)
    assert "Phase 1 : Unique - 1 JEH" in text


def test_saved_file_reopens_and_every_table_row_matches_its_grid(doc):
    for tbl in doc.element.body.iter(qn("w:tbl")):
        grid = len(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")))
        for tr in tbl.findall(qn("w:tr")):
            cols = 0
            for tc in tr.findall(qn("w:tc")):
                span = tc.find(qn("w:tcPr")).find(qn("w:gridSpan")) if tc.find(qn("w:tcPr")) is not None else None
                cols += int(span.get(qn("w:val"))) if span is not None else 1
            assert cols == grid, f"ligne de {cols} colonnes dans une grille de {grid}"


def test_budget_total_is_the_sum_of_the_phase_budgets(doc):
    rows = _rows(_table(doc, "Prix H.T. / JEH"))
    flat = [" ".join(r) for r in rows]
    assert not any("Frais de structure" in r for r in flat)
    phase_totals = [r[-1] for r in rows if r[0].startswith("Total phase")]
    assert phase_totals == ["361,50 € HT", "964,00 € HT", "241,00 € HT"]
    total = next(r for r in rows if r[0].startswith("Total H.T.") and "hors" not in r[0])
    assert any("1 566,50 € HT" in cell for cell in total)
    assert "frais_lettres" not in _text(doc)


def test_text_pasted_from_word_cannot_break_the_generation():
    doc = _render(contexte="Ligne\x0bavec retour Word\x0cet\x00nul\x1b\r\n- puce​zéro﻿￾")
    text = _text(doc)
    assert "Ligne" in text and "avec retour Word" in text and "puce" in text
    assert "\x00" not in text and "​" not in text


def test_president_defaults_to_the_global_variable(doc):
    text = _text(doc)
    assert "César Colombo" in text
    assert "{" not in text


def test_braces_typed_by_the_user_stay_literal():
    doc = _render(cdc="Le montant {total_ht} et {nom_societe} restent tels quels", client__societe="Acme {x}")
    text = _text(doc)
    assert "{total_ht} et {nom_societe} restent tels quels" in text
    assert "Acme {x}" in text
    assert "" not in text and "" not in text
